"""生成 PPTX / 图片解析样例（仅用于测试）。"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "samples"


def make_pptx() -> None:
    presentation = Presentation()
    slides = [
        (
            "西部计划政策要点",
            ["服务期：1 至 3 年，协议一年一签", "服务期满考核合格颁发志愿服务证书"],
        ),
        (
            "补贴与社保",
            ["中央财政给予生活补贴", "服务地可给予配套补贴", "按规定参加社会保险"],
        ),
        (
            "期满优惠",
            ["考研、公务员招录按规定享受优惠", "具体以当年政策文件为准"],
        ),
    ]
    for title, bullets in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for line in bullets[1:]:
            paragraph = body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(20)
        slide.notes_slide.notes_text_frame.text = "本页仅用于解析测试。"
    presentation.save(SAMPLES / "policy_slides.pptx")


def make_png() -> None:
    width, height = 1000, 420
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    font = (
        ImageFont.truetype(str(font_path), 34)
        if font_path.exists()
        else ImageFont.load_default()
    )
    lines = [
        "西部计划政策测试图片",
        "志愿者服务期一般为 1 至 3 年。",
        "服务期间按规定参加社会保险。",
        "生活补贴标准以当年文件为准。",
    ]
    for index, line in enumerate(lines):
        draw.text((40, 40 + index * 85), line, fill=(20, 20, 30), font=font)
    image.save(SAMPLES / "ocr_policy.png")


if __name__ == "__main__":
    SAMPLES.mkdir(parents=True, exist_ok=True)
    make_pptx()
    make_png()
    print("samples generated")
